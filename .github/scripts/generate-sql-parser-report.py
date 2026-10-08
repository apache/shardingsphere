#!/usr/bin/env python3
#
# Licensed to the Apache Software Foundation (ASF) under one or more
# contributor license agreements.  See the NOTICE file distributed with
# this work for additional information regarding copyright ownership.
# The ASF licenses this file to You under the Apache License, Version 2.0
# (the "License"); you may not use this file except in compliance with
# the License.  You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
#

import argparse
from collections import Counter
from pathlib import Path
import re
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill


def split_cell_text(value):
    if len(value.encode("utf-16-le")) // 2 <= 32767 and value.count("\n") <= 253:
        return [value]
    chunks = []
    start = units = line_feeds = 0
    for offset, character in enumerate(value):
        size = 2 if ord(character) > 0xFFFF else 1
        if units + size > 32767 or line_feeds + (character == "\n") > 253:
            chunks.append(value[start:offset])
            start, units, line_feeds = offset, 0, 0
        units += size
        line_feeds += character == "\n"
    chunks.append(value[start:])
    return chunks


def append_row(sheet, values):
    cells = []
    for value in values:
        cell = Cell(sheet, value=value)
        if isinstance(value, str):
            cell.data_type = "s"
        cells.append(cell)
    sheet.append(cells)


def append_error(details, long_text, values):
    for index, value in enumerate(values):
        if not isinstance(value, str):
            continue
        chunks = split_cell_text(value)
        if len(chunks) == 1:
            continue
        field = details.cell(1, index + 1).value
        for part, chunk in enumerate(chunks, 1):
            append_row(long_text, [values[0], field, part, chunk])
        values[index] = f"See Long Text: error {values[0]}, field '{field}' ({len(chunks)} parts)."
    append_row(details, values)


def format_sheet(sheet, widths):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 60


def generate_report(reports_dir, test_class, database, output):
    reports = sorted(reports_dir.glob(f"TEST-*{test_class}.xml"))
    if not reports:
        raise FileNotFoundError(f"No Failsafe XML reports for {test_class} in {reports_dir}")
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    details = workbook.create_sheet("Errors")
    long_text = workbook.create_sheet("Long Text")
    append_row(details, ["Error ID", "Database", "Case ID", "SQL", "Result", "Exception Type", "Error Message", "Stack Trace", "Test Name", "Report File"])
    append_row(long_text, ["Error ID", "Field", "Part", "Text"])
    totals = Counter()
    exceptions = Counter()
    for report in reports:
        for _, testcase in ET.iterparse(report, events=("end",)):
            if testcase.tag != "testcase":
                continue
            totals["tests"] += 1
            totals["skipped"] += testcase.find("skipped") is not None
            totals["errors"] += testcase.find("error") is not None
            totals["failures"] += testcase.find("failure") is not None
            name = testcase.attrib["name"]
            parameters = re.fullmatch(r"assertParseSQL\([^)]*\) (.*?) \(([^()]*)\) -> (.*)", name, re.DOTALL)
            for failure in list(testcase.findall("error")) + list(testcase.findall("failure")):
                totals["error_entries"] += 1
                exceptions[failure.get("type", "")] += 1
                append_error(details, long_text, [
                    totals["error_entries"], parameters[2] if parameters else database, parameters[1] if parameters else "", parameters[3] if parameters else "",
                    failure.tag, failure.get("type", ""), failure.get("message", ""), failure.text or "", name, report.name,
                ])
            testcase.clear()
    append_row(summary, ["Metric", "Count or Value"])
    for label, value in [
        ("Database", database), ("Test Class", test_class), ("Total Tests", totals["tests"]),
        ("Errors", totals["errors"]), ("Failures", totals["failures"]), ("Skipped", totals["skipped"]),
        ("Passed", totals["tests"] - totals["errors"] - totals["failures"] - totals["skipped"]),
        ("Notes", "Only error/failure entries are listed. Initialization failures may have no SQL. Text exceeding Excel cell limits is stored in Long Text by error ID and field."),
    ]:
        append_row(summary, [label, value])
    append_row(summary, ["Exception Type", "Error Entries"])
    for exception, count in exceptions.most_common():
        append_row(summary, [exception, count])
    format_sheet(summary, {"A": 65, "B": 100})
    format_sheet(details, {"A": 12, "B": 16, "C": 45, "D": 100, "E": 12, "F": 65, "G": 100, "H": 100, "I": 65, "J": 65})
    format_sheet(long_text, {"A": 12, "B": 20, "C": 12, "D": 120})
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"Excel report: {output}; tests={totals['tests']}; errors={totals['errors']}; failures={totals['failures']}")


def main():
    parser = argparse.ArgumentParser(description="Export SQL parser Failsafe errors to Excel.")
    parser.add_argument("--reports-dir", type=Path, required=True)
    parser.add_argument("--test-class", required=True)
    parser.add_argument("--database", required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    generate_report(args.reports_dir, args.test_class, args.database, args.output)


if __name__ == "__main__":
    main()
